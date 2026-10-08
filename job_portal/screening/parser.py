"""Resume text extraction and field extraction (Sections 3.7.1 and 3.7.2 of the report)."""

import re
from pathlib import Path

from .skills import SYNONYM_TO_SKILL

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
# Indian mobile number: optional +91 / 91 / 0 prefix, then 10 digits starting 6-9
PHONE_RE = re.compile(r"(?:\+?91[\s-]?|0)?([6-9]\d{4}[\s-]?\d{5})\b")
YEARS_RE = re.compile(r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\b", re.IGNORECASE)

# Highest level wins: diploma 1, bachelor 2, master/MBA 3, doctorate 4.
EDUCATION_LEVELS = [
    (4, r"ph\.?\s?d|doctorate|doctor of philosophy"),
    (3, r"m\.?\s?tech|m\.e\.|m\.?\s?sc|m\.?\s?com|mca|mba|pgdm|master'?s?|post[- ]graduate"),
    (2, r"b\.?\s?tech|b\.e\.|b\.?\s?sc|b\.?\s?com|bca|bba|b\.a\.|bachelor'?s?|graduate|degree"),
    (1, r"diploma|polytechnic|12th|hsc"),
]
EDUCATION_NAMES = {0: "Not specified", 1: "Diploma", 2: "Bachelor's", 3: "Master's / MBA", 4: "Doctorate"}


def extract_text(path):
    """Read plain text from a PDF, DOCX or TXT resume."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == ".docx":
        import docx
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {ext}")


def normalise(text):
    """Lower-case, strip special characters and collapse repeated spaces."""
    text = text.lower()
    text = re.sub(r"[^a-z0-9+#/.\-\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def extract_skills(text):
    """Map every known synonym found in the text to its canonical skill."""
    clean = f" {normalise(text)} "
    found = set()
    for syn, canon in SYNONYM_TO_SKILL.items():
        if re.search(rf"(?<![a-z0-9]){re.escape(syn)}(?![a-z0-9])", clean):
            found.add(canon)
    return found


def extract_experience(text):
    """Largest "N years" figure in the text, taken as years of experience."""
    values = [float(v) for v in YEARS_RE.findall(text)]
    values = [v for v in values if v <= 45]  # ignore obvious noise such as "100 years"
    return max(values) if values else 0.0


def extract_education(text):
    """Highest education level on the 0-4 scale."""
    low = text.lower()
    for level, pattern in EDUCATION_LEVELS:
        if re.search(rf"\b(?:{pattern})", low):
            return level
    return 0


def extract_email(text):
    match = EMAIL_RE.search(text)
    return match.group(0) if match else None


def extract_phone(text):
    match = PHONE_RE.search(text)
    return re.sub(r"[\s-]", "", match.group(1)) if match else None


def parse_resume(text):
    return {
        "clean_text": normalise(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "experience_years": extract_experience(text),
        "education_level": extract_education(text),
        "skills": extract_skills(text),
    }


def parse_job(description, min_experience=None, min_education=None, skills=None):
    """Parse a job description. Explicit values override what is read from the text."""
    return {
        "clean_text": normalise(description),
        "skills": set(skills) if skills else extract_skills(description),
        "min_experience": float(min_experience) if min_experience is not None
        else extract_experience(description),
        "min_education": int(min_education) if min_education is not None
        else extract_education(description),
    }
