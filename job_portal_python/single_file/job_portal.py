"""
===============================================================================
 JobMatch AI - Full-Stack Job Portal with AI-Based Resume Screening
               and Candidate Matching  (single-file Python version)
===============================================================================
 SRM University - Final Project

 HOW TO RUN (VS Code / any terminal):
   1. Install Python 3.8+ (tick "Add Python to PATH" on Windows).
   2. Save this file as  job_portal.py  in an empty folder.
   3. Open a terminal in that folder and run:
          pip install flask flask-sqlalchemy flask-login scikit-learn pypdf python-docx
   4. Start the app:
          python job_portal.py
   5. Open  http://127.0.0.1:5000  in your browser.

 Demo data is created automatically on first run. All demo passwords: password123
     Admin     : admin@jobmatch.ai
     Employer  : hr@techsoft.com      talent@cloudnine.io
     Candidate : priya@example.com    rahul@example.com
                 ananya@example.com   arjun@example.com
 To reset the demo data, delete the file  jobportal.db  and run again.

 TECH STACK
   Backend  : Flask (Python web framework), Flask-Login (sessions/auth)
   Database : SQLite via Flask-SQLAlchemy (ORM)
   AI / NLP : scikit-learn TF-IDF + cosine similarity, regex-based information
              extraction (skills, experience, education)
   Parsing  : pypdf (PDF), python-docx (DOCX)
   Frontend : HTML (Jinja2 templates) + Bootstrap 5

 AI MATCH SCORE (0-100) = 50% skill match + 30% TF-IDF text similarity
                          + 20% experience fit
   (if the resume does not mention years of experience, that 20% is
    redistributed over the other two so the candidate is not penalised)
===============================================================================
"""

import io
import json
import os
import re
import secrets
import uuid
from datetime import datetime
from functools import wraps

from flask import (Flask, abort, flash, redirect, render_template, request,
                   send_from_directory, session, url_for)
from flask_login import (LoginManager, UserMixin, current_user, login_required,
                         login_user, logout_user)
from flask_sqlalchemy import SQLAlchemy
from jinja2 import DictLoader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from sqlalchemy import func
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

# =============================================================================
# 1. CONFIGURATION
# =============================================================================
BASE_DIR = os.path.abspath(os.path.dirname(__file__))
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-this-secret-key")
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + os.path.join(BASE_DIR, "jobportal.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024  # 5 MB upload limit
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"
login_manager.login_message_category = "warning"

JOB_TYPES = ["Full-time", "Part-time", "Internship", "Contract", "Remote"]
STATUSES = ["applied", "shortlisted", "interview", "rejected", "hired"]


# =============================================================================
# 2. DATABASE MODELS
# =============================================================================
class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=False)
    role = db.Column(db.String(20), default="candidate")  # candidate / employer / admin
    active_account = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    # candidate profile
    headline = db.Column(db.String(200))
    location = db.Column(db.String(120))
    skills = db.Column(db.Text, default="")
    resume_filename = db.Column(db.String(255))
    resume_text = db.Column(db.Text)
    # employer profile
    company = db.Column(db.String(150))

    jobs = db.relationship("Job", backref="employer", cascade="all, delete-orphan")
    applications = db.relationship("Application", backref="candidate", cascade="all, delete-orphan")

    def set_password(self, pw):
        self.password_hash = generate_password_hash(pw)

    def check_password(self, pw):
        return check_password_hash(self.password_hash, pw)

    @property
    def is_active(self):
        return bool(self.active_account)

    @property
    def skill_list(self):
        return [s.strip() for s in (self.skills or "").split(",") if s.strip()]


class Job(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(150), nullable=False)
    company = db.Column(db.String(150), nullable=False)
    location = db.Column(db.String(120), nullable=False)
    job_type = db.Column(db.String(30), default="Full-time")
    salary = db.Column(db.String(60))
    min_experience = db.Column(db.Integer, default=0)
    required_skills = db.Column(db.Text, nullable=False)
    description = db.Column(db.Text, nullable=False)
    is_open = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    employer_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    applications = db.relationship("Application", backref="job", cascade="all, delete-orphan")

    @property
    def skill_list(self):
        return [s.strip() for s in (self.required_skills or "").split(",") if s.strip()]


class Application(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    job_id = db.Column(db.Integer, db.ForeignKey("job.id"), nullable=False)
    candidate_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    cover_letter = db.Column(db.Text)
    resume_filename = db.Column(db.String(255))
    resume_text = db.Column(db.Text)
    status = db.Column(db.String(20), default="applied")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    # AI screening results
    match_score = db.Column(db.Float, default=0)
    skill_score = db.Column(db.Float)
    text_score = db.Column(db.Float)
    experience_score = db.Column(db.Float)
    years_experience = db.Column(db.Float)
    education = db.Column(db.String(50))
    matched_skills = db.Column(db.Text, default="[]")
    missing_skills = db.Column(db.Text, default="[]")
    recommendation = db.Column(db.String(30))
    __table_args__ = (db.UniqueConstraint("job_id", "candidate_id"),)

    @property
    def matched_list(self):
        return json.loads(self.matched_skills or "[]")

    @property
    def missing_list(self):
        return json.loads(self.missing_skills or "[]")

    def apply_screening(self, r):
        self.match_score, self.skill_score, self.text_score = r["score"], r["skill_score"], r["text_score"]
        self.experience_score, self.years_experience = r["experience_score"], r["years"]
        self.education, self.recommendation = r["education"], r["recommendation"]
        self.matched_skills, self.missing_skills = json.dumps(r["matched"]), json.dumps(r["missing"])


@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))


# =============================================================================
# 3. AI ENGINE - RESUME PARSING
# =============================================================================
def extract_text(file_bytes, filename):
    """Convert an uploaded resume (PDF / DOCX / TXT) into plain text."""
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == "docx":
        import docx
        document = docx.Document(io.BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if ext == "txt":
        return file_bytes.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type")


# =============================================================================
# 4. AI ENGINE - INFORMATION EXTRACTION (skills, experience, education)
# =============================================================================
SKILLS = [
    "python", "java", "javascript", "typescript", "c", "c++", "c#", "go", "rust", "kotlin",
    "swift", "php", "ruby", "r", "scala", "matlab", "sql", "bash",
    "html", "css", "react", "angular", "vue", "node.js", "express", "django", "flask",
    "fastapi", "spring boot", "bootstrap", "tailwind", "rest api", "graphql", "jquery",
    "machine learning", "deep learning", "nlp", "computer vision", "data analysis",
    "data science", "statistics", "pandas", "numpy", "scikit-learn", "tensorflow",
    "pytorch", "keras", "opencv", "matplotlib", "power bi", "tableau", "excel",
    "big data", "spark", "hadoop", "llm", "generative ai",
    "mysql", "postgresql", "mongodb", "sqlite", "oracle", "redis", "firebase",
    "aws", "azure", "gcp", "docker", "kubernetes", "jenkins", "ci/cd", "git", "linux",
    "terraform", "microservices", "android", "ios", "flutter", "react native",
    "selenium", "unit testing", "agile", "scrum", "jira", "cyber security", "networking",
    "communication", "leadership", "teamwork", "problem solving", "project management",
]

# Synonyms / abbreviations  ->  canonical skill name
ALIASES = {
    "js": "javascript", "es6": "javascript", "ts": "typescript", "reactjs": "react",
    "react.js": "react", "nodejs": "node.js", "expressjs": "express", "vuejs": "vue",
    "angularjs": "angular", "golang": "go", "cpp": "c++", "csharp": "c#",
    "ml": "machine learning", "dl": "deep learning", "natural language processing": "nlp",
    "sklearn": "scikit-learn", "scikit learn": "scikit-learn", "tf": "tensorflow",
    "postgres": "postgresql", "mongo": "mongodb", "amazon web services": "aws",
    "google cloud": "gcp", "microsoft azure": "azure", "k8s": "kubernetes",
    "restful": "rest api", "rest apis": "rest api", "powerbi": "power bi",
    "ms excel": "excel", "springboot": "spring boot", "cybersecurity": "cyber security",
    "large language models": "llm", "gen ai": "generative ai", "html5": "html",
    "css3": "css", "github": "git", "unit tests": "unit testing",
}

# Knowing a specific tool implies knowing the broader skill
IMPLIES = {
    "mysql": ["sql"], "postgresql": ["sql"], "sqlite": ["sql"], "oracle": ["sql"],
    "django": ["python"], "flask": ["python"], "fastapi": ["python"], "pandas": ["python"],
    "react": ["javascript"], "angular": ["javascript"], "vue": ["javascript"],
    "node.js": ["javascript"], "express": ["javascript", "node.js"], "typescript": ["javascript"],
    "spring boot": ["java"], "tensorflow": ["deep learning"], "pytorch": ["deep learning"],
    "keras": ["deep learning"], "deep learning": ["machine learning"],
    "scikit-learn": ["machine learning"], "kubernetes": ["docker"],
}

# Short names that are also English words: only counted inside a list ("Python, C, Go")
AMBIGUOUS = {"c", "r", "go"}

EDUCATION_LEVELS = [
    ("PhD", r"\b(ph\.?\s?d|doctorate)\b"),
    ("Master's", r"\b(m\.?\s?tech\b|m\.e\.|m\.?\s?sc\b|mca\b|mba\b|master'?s?\b|m\.s\.)"),
    ("Bachelor's", r"\b(b\.?\s?tech\b|b\.e\.|b\.?\s?sc\b|bca\b|bba\b|bachelor'?s?\b|b\.s\.)"),
    ("Diploma", r"\bdiploma\b"),
    ("High School", r"\b(high school|12th|hsc|higher secondary)\b"),
]


def _pattern(term):
    # boundaries that also work for "c++", "c#", "node.js"
    return re.compile(r"(?<![a-z0-9+#])" + re.escape(term) + r"(?![a-z0-9+#])")


TERM_PATTERNS = [(t, _pattern(t), t) for t in SKILLS] + \
                [(a, _pattern(a), c) for a, c in ALIASES.items()]


def normalize_skill(skill):
    s = skill.strip().lower()
    return ALIASES.get(s, s)


def expand_implied(skills):
    result, pending = set(skills), list(skills)
    while pending:
        for implied in IMPLIES.get(pending.pop(), []):
            if implied not in result:
                result.add(implied)
                pending.append(implied)
    return result


def extract_skills(text):
    if not text:
        return []
    lowered, found = text.lower(), set()
    for term, pattern, canonical in TERM_PATTERNS:
        if term in AMBIGUOUS and not re.search(
                r"(?:^|[,:/|(]\s*)" + re.escape(term) + r"\s*(?:[,/|)]|$)", lowered, re.M):
            continue
        if pattern.search(lowered):
            found.add(canonical)
    return sorted(expand_implied(found))


def extract_years_experience(text):
    if not text:
        return None
    matches = re.findall(
        r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:\w+\s+){0,3}?(?:experience|exp)",
        text.lower())
    years = [float(m) for m in matches if float(m) <= 50]
    return max(years) if years else None


def extract_education(text):
    if not text:
        return None
    for level, pattern in EDUCATION_LEVELS:
        if re.search(pattern, text.lower()):
            return level
    return None


# =============================================================================
# 5. AI ENGINE - SCREENING, CANDIDATE RANKING & JOB RECOMMENDATION
# =============================================================================
WEIGHTS = {"skill": 0.5, "text": 0.3, "experience": 0.2}
SIMILARITY_CEILING = 0.5  # cosine >= 0.5 between a resume and a job counts as 100%


def text_similarity(resume_text, job_text):
    """TF-IDF vectors + cosine similarity (0..1)."""
    if not (resume_text or "").strip() or not (job_text or "").strip():
        return 0.0
    vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
    try:
        matrix = vectorizer.fit_transform([resume_text, job_text])
    except ValueError:
        return 0.0
    return min(1.0, max(0.0, float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0])))


def recommendation_for(score):
    if score >= 75:
        return "Strong Match"
    if score >= 50:
        return "Good Match"
    if score >= 30:
        return "Partial Match"
    return "Low Match"


def score_resume(resume_text, job, profile_skills=None):
    resume_text = resume_text or ""
    resume_skills = set(extract_skills(resume_text))
    resume_skills = expand_implied(resume_skills | {normalize_skill(s) for s in (profile_skills or [])})

    job_skills = []
    for s in job.skill_list:
        if normalize_skill(s) not in job_skills:
            job_skills.append(normalize_skill(s))
    matched = [s for s in job_skills if s in resume_skills]
    missing = [s for s in job_skills if s not in resume_skills]

    # (1) skill match
    skill_score = len(matched) / len(job_skills) * 100 if job_skills else 100.0
    # (2) TF-IDF cosine similarity between resume and job text
    job_text = "%s. %s. Skills: %s" % (job.title, job.description, job.required_skills)
    text_score = min(1.0, text_similarity(resume_text, job_text) / SIMILARITY_CEILING) * 100
    # (3) experience fit
    years = extract_years_experience(resume_text)
    if not job.min_experience:
        exp_score = 100.0
    elif years is None:
        exp_score = None
    else:
        exp_score = round(min(1.0, years / job.min_experience) * 100, 1)

    if exp_score is None:  # redistribute the experience weight
        score = (WEIGHTS["skill"] * skill_score + WEIGHTS["text"] * text_score) / (WEIGHTS["skill"] + WEIGHTS["text"])
    else:
        score = WEIGHTS["skill"] * skill_score + WEIGHTS["text"] * text_score + WEIGHTS["experience"] * exp_score
    score = round(score, 1)

    return {
        "score": score, "skill_score": round(skill_score, 1), "text_score": round(text_score, 1),
        "experience_score": exp_score, "years": years, "education": extract_education(resume_text),
        "matched": matched, "missing": missing, "recommendation": recommendation_for(score),
    }


def rank_candidates(job, candidates):
    """Talent search: rank every candidate profile for one job."""
    results = [(c, score_resume(c.resume_text, job, c.skill_list))
               for c in candidates if c.resume_text or c.skill_list]
    return sorted(results, key=lambda x: x[1]["score"], reverse=True)


def recommend_jobs(candidate, jobs, limit=5):
    """Job recommendation: rank open jobs for one candidate."""
    results = [(j, score_resume(candidate.resume_text, j, candidate.skill_list)) for j in jobs]
    return sorted(results, key=lambda x: x[1]["score"], reverse=True)[:limit]


# =============================================================================
# 6. HELPERS - roles, CSRF protection, resume upload
# =============================================================================
def role_required(*roles):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not current_user.is_authenticated:
                return login_manager.unauthorized()
            if current_user.role not in roles:
                abort(403)
            return view(*args, **kwargs)
        return wrapped
    return decorator


def csrf_token():
    if "_csrf" not in session:
        session["_csrf"] = secrets.token_hex(16)
    return session["_csrf"]


app.jinja_env.globals["csrf_token"] = csrf_token


@app.before_request
def csrf_protect():
    if request.method == "POST" and not app.config.get("TESTING"):
        if request.form.get("csrf_token") != session.get("_csrf"):
            abort(400)


def save_resume(file_storage):
    """Validate + store + parse a resume. Returns (stored_filename, text)."""
    if not file_storage or not file_storage.filename:
        raise ValueError("Please choose a resume file.")
    if file_storage.filename.rsplit(".", 1)[-1].lower() not in ALLOWED_EXTENSIONS:
        raise ValueError("Resume must be a PDF, DOCX or TXT file.")
    data = file_storage.read()
    try:
        text = extract_text(data, file_storage.filename)
    except Exception:
        raise ValueError("Could not read that resume file.")
    if not text.strip():
        raise ValueError("No text found in the resume (scanned images are not supported).")
    stored = uuid.uuid4().hex + "_" + secure_filename(file_storage.filename)
    with open(os.path.join(UPLOAD_FOLDER, stored), "wb") as fh:
        fh.write(data)
    return stored, text


def home_for(user):
    return url_for({"employer": "employer_dashboard", "admin": "admin_dashboard"}.get(
        user.role, "candidate_dashboard"))


def own_job_or_403(job_id):
    job = db.get_or_404(Job, job_id)
    if job.employer_id != current_user.id and current_user.role != "admin":
        abort(403)
    return job


@app.template_filter("color")
def score_color(score):
    if score is None:
        return "secondary"
    return "success" if score >= 75 else "primary" if score >= 50 else "warning" if score >= 30 else "danger"


# =============================================================================
# 7. ROUTES - PUBLIC & AUTHENTICATION
# =============================================================================
@app.route("/")
def index():
    stats = {
        "jobs": Job.query.filter_by(is_open=True).count(),
        "companies": db.session.query(Job.company).distinct().count(),
        "candidates": User.query.filter_by(role="candidate").count(),
        "applications": Application.query.count(),
    }
    jobs = Job.query.filter_by(is_open=True).order_by(Job.created_at.desc()).limit(6).all()
    return render_template("index.html", stats=stats, jobs=jobs)


@app.route("/register", methods=["GET", "POST"])
def register():
    if current_user.is_authenticated:
        return redirect(home_for(current_user))
    if request.method == "POST":
        f = request.form
        name, email, pw = f.get("name", "").strip(), f.get("email", "").strip().lower(), f.get("password", "")
        role, company = f.get("role", "candidate"), f.get("company", "").strip()
        error = None
        if not name or not email or not pw:
            error = "All fields are required."
        elif len(pw) < 6:
            error = "Password must be at least 6 characters."
        elif role not in ("candidate", "employer"):
            error = "Invalid role."
        elif role == "employer" and not company:
            error = "Company name is required for employers."
        elif User.query.filter_by(email=email).first():
            error = "An account with that email already exists."
        if error:
            flash(error, "danger")
        else:
            user = User(name=name, email=email, role=role, company=company or None)
            user.set_password(pw)
            db.session.add(user)
            db.session.commit()
            login_user(user)
            flash("Welcome, %s!" % name, "success")
            return redirect(home_for(user))
    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(home_for(current_user))
    if request.method == "POST":
        user = User.query.filter_by(email=request.form.get("email", "").strip().lower()).first()
        if not user or not user.check_password(request.form.get("password", "")):
            flash("Invalid email or password.", "danger")
        elif not user.is_active:
            flash("This account has been deactivated.", "danger")
        else:
            login_user(user)
            nxt = request.args.get("next")
            if nxt and nxt.startswith("/") and not nxt.startswith("//"):
                return redirect(nxt)
            return redirect(home_for(user))
    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    flash("You have been logged out.", "info")
    return redirect(url_for("index"))


@app.route("/jobs")
def jobs():
    q, loc, jt = request.args.get("q", "").strip(), request.args.get("location", "").strip(), request.args.get("job_type", "")
    query = Job.query.filter_by(is_open=True)
    if q:
        like = "%" + q + "%"
        query = query.filter(Job.title.ilike(like) | Job.description.ilike(like) |
                             Job.required_skills.ilike(like) | Job.company.ilike(like))
    if loc:
        query = query.filter(Job.location.ilike("%" + loc + "%"))
    if jt in JOB_TYPES:
        query = query.filter_by(job_type=jt)
    return render_template("jobs.html", jobs=query.order_by(Job.created_at.desc()).all(),
                           q=q, location=loc, job_type=jt, job_types=JOB_TYPES)


@app.route("/jobs/<int:job_id>")
def job_detail(job_id):
    job = db.get_or_404(Job, job_id)
    existing = preview = None
    if current_user.is_authenticated and current_user.role == "candidate":
        existing = Application.query.filter_by(job_id=job.id, candidate_id=current_user.id).first()
        if not existing and (current_user.resume_text or current_user.skill_list):
            preview = score_resume(current_user.resume_text, job, current_user.skill_list)
    return render_template("job_detail.html", job=job, existing=existing, preview=preview)


@app.route("/jobs/<int:job_id>/apply", methods=["POST"])
@role_required("candidate")
def apply(job_id):
    job = db.get_or_404(Job, job_id)
    back = redirect(url_for("job_detail", job_id=job.id))
    if not job.is_open:
        flash("This job is closed.", "warning")
        return back
    if Application.query.filter_by(job_id=job.id, candidate_id=current_user.id).first():
        flash("You have already applied to this job.", "info")
        return back
    upload = request.files.get("resume")
    if upload and upload.filename:
        try:
            stored, text = save_resume(upload)
        except ValueError as e:
            flash(str(e), "danger")
            return back
        current_user.resume_filename, current_user.resume_text = stored, text
    elif current_user.resume_text:
        stored, text = current_user.resume_filename, current_user.resume_text
    else:
        flash("Please upload a resume to apply.", "danger")
        return back

    application = Application(job_id=job.id, candidate_id=current_user.id, resume_filename=stored,
                              resume_text=text, cover_letter=request.form.get("cover_letter", ""))
    result = score_resume(text, job, current_user.skill_list)  # <-- AI screening
    application.apply_screening(result)
    db.session.add(application)
    db.session.commit()
    flash("Application submitted! AI match score: %.0f%% (%s)" % (result["score"], result["recommendation"]), "success")
    return redirect(url_for("candidate_application", app_id=application.id))


# =============================================================================
# 8. ROUTES - CANDIDATE
# =============================================================================
@app.route("/candidate")
@role_required("candidate")
def candidate_dashboard():
    apps = Application.query.filter_by(candidate_id=current_user.id).order_by(Application.created_at.desc()).all()
    applied = {a.job_id for a in apps}
    recs = []
    if current_user.resume_text or current_user.skill_list:
        open_jobs = [j for j in Job.query.filter_by(is_open=True).all() if j.id not in applied]
        recs = recommend_jobs(current_user, open_jobs)
    return render_template("candidate_dashboard.html", applications=apps, recommendations=recs)


@app.route("/candidate/profile", methods=["GET", "POST"])
@role_required("candidate")
def candidate_profile():
    if request.method == "POST":
        f = request.form
        current_user.name = f.get("name", "").strip() or current_user.name
        current_user.headline = f.get("headline", "").strip()
        current_user.location = f.get("location", "").strip()
        current_user.skills = ", ".join(s.strip() for s in f.get("skills", "").split(",") if s.strip())
        upload = request.files.get("resume")
        if upload and upload.filename:
            try:
                current_user.resume_filename, current_user.resume_text = save_resume(upload)
            except ValueError as e:
                flash(str(e), "danger")
                return redirect(url_for("candidate_profile"))
        db.session.commit()
        flash("Profile updated.", "success")
        return redirect(url_for("candidate_profile"))
    analysis = None
    if current_user.resume_text:
        t = current_user.resume_text
        analysis = {"skills": extract_skills(t), "years": extract_years_experience(t),
                    "education": extract_education(t), "words": len(t.split())}
    return render_template("candidate_profile.html", analysis=analysis)


@app.route("/candidate/applications/<int:app_id>")
@role_required("candidate")
def candidate_application(app_id):
    a = db.get_or_404(Application, app_id)
    if a.candidate_id != current_user.id:
        abort(403)
    return render_template("report.html", a=a, employer_view=False)


@app.route("/candidate/applications/<int:app_id>/withdraw", methods=["POST"])
@role_required("candidate")
def withdraw(app_id):
    a = db.get_or_404(Application, app_id)
    if a.candidate_id != current_user.id:
        abort(403)
    db.session.delete(a)
    db.session.commit()
    flash("Application withdrawn.", "info")
    return redirect(url_for("candidate_dashboard"))


# =============================================================================
# 9. ROUTES - EMPLOYER
# =============================================================================
def job_from_form(job):
    f = request.form
    job.title, job.location = f.get("title", "").strip(), f.get("location", "").strip()
    job.company = f.get("company", "").strip() or current_user.company or ""
    job.job_type = f.get("job_type") if f.get("job_type") in JOB_TYPES else "Full-time"
    job.salary, job.description = f.get("salary", "").strip(), f.get("description", "").strip()
    try:
        job.min_experience = max(0, int(f.get("min_experience") or 0))
    except ValueError:
        job.min_experience = 0
    job.required_skills = ", ".join(s.strip() for s in f.get("required_skills", "").split(",") if s.strip())
    if not all([job.title, job.company, job.location, job.required_skills, job.description]):
        return "Please fill in all required fields."
    return None


@app.route("/employer")
@role_required("employer")
def employer_dashboard():
    my_jobs = Job.query.filter_by(employer_id=current_user.id).order_by(Job.created_at.desc()).all()
    total = sum(len(j.applications) for j in my_jobs)
    shortlisted = sum(1 for j in my_jobs for a in j.applications if a.status == "shortlisted")
    return render_template("employer_dashboard.html", jobs=my_jobs, total=total, shortlisted=shortlisted)


@app.route("/employer/jobs/new", methods=["GET", "POST"])
@role_required("employer")
def new_job():
    job = Job(company=current_user.company, min_experience=0)
    if request.method == "POST":
        error = job_from_form(job)
        if error:
            flash(error, "danger")
        else:
            job.employer_id = current_user.id
            db.session.add(job)
            db.session.commit()
            flash("Job posted successfully.", "success")
            return redirect(url_for("employer_dashboard"))
    return render_template("job_form.html", job=job, job_types=JOB_TYPES, editing=False)


@app.route("/employer/jobs/<int:job_id>/edit", methods=["GET", "POST"])
@role_required("employer", "admin")
def edit_job(job_id):
    job = own_job_or_403(job_id)
    if request.method == "POST":
        error = job_from_form(job)
        if error:
            flash(error, "danger")
        else:
            for a in job.applications:  # requirements changed -> re-screen everyone
                a.apply_screening(score_resume(a.resume_text, job, a.candidate.skill_list))
            db.session.commit()
            flash("Job updated and applicants re-screened.", "success")
            return redirect(url_for("applicants", job_id=job.id))
    return render_template("job_form.html", job=job, job_types=JOB_TYPES, editing=True)


@app.route("/employer/jobs/<int:job_id>/toggle", methods=["POST"])
@role_required("employer", "admin")
def toggle_job(job_id):
    job = own_job_or_403(job_id)
    job.is_open = not job.is_open
    db.session.commit()
    flash("Job %s." % ("reopened" if job.is_open else "closed"), "info")
    return redirect(request.referrer or url_for("employer_dashboard"))


@app.route("/employer/jobs/<int:job_id>/delete", methods=["POST"])
@role_required("employer", "admin")
def delete_job(job_id):
    db.session.delete(own_job_or_403(job_id))
    db.session.commit()
    flash("Job deleted.", "info")
    return redirect(url_for("admin_jobs" if current_user.role == "admin" else "employer_dashboard"))


@app.route("/employer/jobs/<int:job_id>/applicants")
@role_required("employer", "admin")
def applicants(job_id):
    job = own_job_or_403(job_id)
    status = request.args.get("status", "")
    min_score = request.args.get("min_score", type=float) or 0
    apps = [a for a in job.applications if a.match_score >= min_score and (not status or a.status == status)]
    apps.sort(key=lambda a: a.match_score, reverse=True)  # ranked by AI score
    return render_template("applicants.html", job=job, applications=apps, statuses=STATUSES,
                           status=status, min_score=min_score)


@app.route("/employer/jobs/<int:job_id>/auto-shortlist", methods=["POST"])
@role_required("employer", "admin")
def auto_shortlist(job_id):
    job = own_job_or_403(job_id)
    threshold = request.form.get("threshold", type=float) or 70
    count = 0
    for a in job.applications:
        if a.status == "applied" and a.match_score >= threshold:
            a.status, count = "shortlisted", count + 1
    db.session.commit()
    flash("AI auto-shortlisted %d candidate(s) scoring %.0f%% or higher." % (count, threshold), "success")
    return redirect(url_for("applicants", job_id=job.id))


@app.route("/employer/jobs/<int:job_id>/talent")
@role_required("employer", "admin")
def talent_search(job_id):
    job = own_job_or_403(job_id)
    candidates = User.query.filter_by(role="candidate", active_account=True).all()
    applied = {a.candidate_id for a in job.applications}
    return render_template("talent.html", job=job, ranked=rank_candidates(job, candidates)[:25], applied=applied)


@app.route("/employer/applications/<int:app_id>")
@role_required("employer", "admin")
def employer_application(app_id):
    a = db.get_or_404(Application, app_id)
    own_job_or_403(a.job_id)
    return render_template("report.html", a=a, employer_view=True, statuses=STATUSES)


@app.route("/employer/applications/<int:app_id>/status", methods=["POST"])
@role_required("employer", "admin")
def update_status(app_id):
    a = db.get_or_404(Application, app_id)
    own_job_or_403(a.job_id)
    if request.form.get("status") in STATUSES:
        a.status = request.form["status"]
        db.session.commit()
        flash("%s marked as %s." % (a.candidate.name, a.status), "success")
    return redirect(request.referrer or url_for("applicants", job_id=a.job_id))


@app.route("/resume/<int:app_id>")
@login_required
def download_resume(app_id):
    a = db.get_or_404(Application, app_id)
    if current_user.role == "candidate":
        if a.candidate_id != current_user.id:
            abort(403)
    else:
        own_job_or_403(a.job_id)
    if not a.resume_filename:
        abort(404)
    return send_from_directory(UPLOAD_FOLDER, a.resume_filename, as_attachment=True)


# =============================================================================
# 10. ROUTES - ADMIN
# =============================================================================
@app.route("/admin")
@role_required("admin")
def admin_dashboard():
    stats = {
        "candidates": User.query.filter_by(role="candidate").count(),
        "employers": User.query.filter_by(role="employer").count(),
        "jobs": Job.query.count(), "applications": Application.query.count(),
        "avg": db.session.query(func.avg(Application.match_score)).scalar() or 0,
    }
    by_status = dict(db.session.query(Application.status, func.count(Application.id))
                     .group_by(Application.status).all())
    buckets = {"75-100": 0, "50-74": 0, "30-49": 0, "0-29": 0}
    for (s,) in db.session.query(Application.match_score).all():
        key = "75-100" if s >= 75 else "50-74" if s >= 50 else "30-49" if s >= 30 else "0-29"
        buckets[key] += 1
    recent = Application.query.order_by(Application.created_at.desc()).limit(10).all()
    return render_template("admin_dashboard.html", stats=stats, by_status=by_status,
                           buckets=buckets, recent=recent)


@app.route("/admin/users")
@role_required("admin")
def admin_users():
    role, q = request.args.get("role", ""), request.args.get("q", "").strip()
    query = User.query
    if role:
        query = query.filter_by(role=role)
    if q:
        query = query.filter(User.name.ilike("%" + q + "%") | User.email.ilike("%" + q + "%"))
    return render_template("admin_users.html", users=query.order_by(User.created_at.desc()).all(), role=role, q=q)


@app.route("/admin/users/<int:user_id>/toggle", methods=["POST"])
@role_required("admin")
def toggle_user(user_id):
    user = db.get_or_404(User, user_id)
    if user.id != current_user.id:
        user.active_account = not user.active_account
        db.session.commit()
        flash("%s %s." % (user.name, "activated" if user.active_account else "deactivated"), "info")
    return redirect(url_for("admin_users"))


@app.route("/admin/users/<int:user_id>/delete", methods=["POST"])
@role_required("admin")
def delete_user(user_id):
    user = db.get_or_404(User, user_id)
    if user.id != current_user.id:
        db.session.delete(user)
        db.session.commit()
        flash("User deleted.", "info")
    return redirect(url_for("admin_users"))


@app.route("/admin/jobs")
@role_required("admin")
def admin_jobs():
    return render_template("admin_jobs.html", jobs=Job.query.order_by(Job.created_at.desc()).all())


@app.errorhandler(403)
def forbidden(e):
    return render_template("error.html", code=403, message="You don't have access to this page."), 403


@app.errorhandler(404)
def not_found(e):
    return render_template("error.html", code=404, message="Page not found."), 404


@app.errorhandler(413)
def too_large(e):
    return render_template("error.html", code=413, message="File too large (max 5 MB)."), 413


# =============================================================================
# 11. HTML TEMPLATES (Jinja2 + Bootstrap 5)
# =============================================================================
H = '<input type="hidden" name="csrf_token" value="{{ csrf_token() }}">'  # CSRF field

TEMPLATES = {
"base.html": """<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{% block title %}JobMatch AI{% endblock %} - JobMatch AI</title>
<link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.min.css" rel="stylesheet">
<link href="https://cdn.jsdelivr.net/npm/bootstrap-icons@1.11.3/font/bootstrap-icons.min.css" rel="stylesheet">
<style>
 .bg-brand{background:linear-gradient(90deg,#4338ca,#6d28d9)} .text-brand{color:#4338ca}
 .chip{background:#eef2ff;color:#3730a3} .hero{background:linear-gradient(135deg,#eef2ff,#f5f3ff);border-radius:1rem}
 .stat{border-left:4px solid #4338ca} .big{font-size:3rem;font-weight:700;line-height:1}
 .card-hover:hover{transform:translateY(-3px);box-shadow:0 .5rem 1rem rgba(67,56,202,.15)!important;transition:.15s}
 .resume{white-space:pre-wrap;max-height:400px;overflow:auto;font-size:.85rem;background:#f8f9fa}
</style></head>
<body class="d-flex flex-column min-vh-100">
<nav class="navbar navbar-expand-lg navbar-dark bg-brand shadow-sm"><div class="container">
 <a class="navbar-brand fw-bold" href="{{ url_for('index') }}"><i class="bi bi-cpu"></i> JobMatch AI</a>
 <button class="navbar-toggler" data-bs-toggle="collapse" data-bs-target="#nav"><span class="navbar-toggler-icon"></span></button>
 <div class="collapse navbar-collapse" id="nav"><ul class="navbar-nav me-auto">
  <li class="nav-item"><a class="nav-link" href="{{ url_for('jobs') }}">Browse Jobs</a></li>
  {% if current_user.is_authenticated %}
   {% if current_user.role == 'candidate' %}
    <li class="nav-item"><a class="nav-link" href="{{ url_for('candidate_dashboard') }}">My Dashboard</a></li>
    <li class="nav-item"><a class="nav-link" href="{{ url_for('candidate_profile') }}">Profile &amp; Resume</a></li>
   {% elif current_user.role == 'employer' %}
    <li class="nav-item"><a class="nav-link" href="{{ url_for('employer_dashboard') }}">My Jobs</a></li>
    <li class="nav-item"><a class="nav-link" href="{{ url_for('new_job') }}">Post a Job</a></li>
   {% else %}
    <li class="nav-item"><a class="nav-link" href="{{ url_for('admin_dashboard') }}">Admin</a></li>
    <li class="nav-item"><a class="nav-link" href="{{ url_for('admin_users') }}">Users</a></li>
    <li class="nav-item"><a class="nav-link" href="{{ url_for('admin_jobs') }}">All Jobs</a></li>
   {% endif %}
  {% endif %}</ul>
  <ul class="navbar-nav">{% if current_user.is_authenticated %}
   <li class="nav-item navbar-text me-3 text-white">{{ current_user.name }} <span class="badge bg-light text-dark text-capitalize">{{ current_user.role }}</span></li>
   <li class="nav-item"><a class="btn btn-outline-light btn-sm mt-1" href="{{ url_for('logout') }}">Logout</a></li>
  {% else %}
   <li class="nav-item"><a class="nav-link" href="{{ url_for('login') }}">Login</a></li>
   <li class="nav-item"><a class="btn btn-light btn-sm mt-1 ms-2" href="{{ url_for('register') }}">Sign up</a></li>
  {% endif %}</ul></div></div></nav>
<main class="container my-4 flex-grow-1">
 {% for cat, msg in get_flashed_messages(with_categories=true) %}
  <div class="alert alert-{{ cat }} alert-dismissible fade show">{{ msg }}<button class="btn-close" data-bs-dismiss="alert"></button></div>
 {% endfor %}
 {% block content %}{% endblock %}
</main>
<footer class="bg-light border-top py-3 text-center text-muted small">JobMatch AI - Full-Stack Job Portal with AI-Based Resume Screening &amp; Candidate Matching | Python, Flask &amp; scikit-learn</footer>
<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/js/bootstrap.bundle.min.js"></script>
</body></html>""",

"macros.html": """
{% macro badge(score) %}<span class="badge rounded-pill bg-{{ score|color }} fs-6">{{ '%.0f'|format(score or 0) }}%</span>{% endmacro %}
{% macro bar(label, value, note='not mentioned - weight redistributed') %}
 <div class="mb-2"><div class="d-flex justify-content-between small"><span>{{ label }}</span>
 <span>{% if value is none %}<em class="text-muted">{{ note }}</em>{% else %}{{ '%.0f'|format(value) }}%{% endif %}</span></div>
 <div class="progress" style="height:8px"><div class="progress-bar bg-{{ value|color }}" style="width:{{ value or 0 }}%"></div></div></div>
{% endmacro %}
{% macro breakdown(skill, text, exp) %}{{ bar('Skill match (50%)', skill) }}{{ bar('TF-IDF text similarity (30%)', text) }}{{ bar('Experience fit (20%)', exp) }}{% endmacro %}
{% macro chips(items, cls='bg-success') %}{% for s in items %}<span class="badge {{ cls }} me-1 mb-1">{{ s }}</span>{% else %}<span class="text-muted small">none</span>{% endfor %}{% endmacro %}
{% macro status(s) %}{% set c = {'applied':'secondary','shortlisted':'info','interview':'primary','rejected':'danger','hired':'success'} %}<span class="badge bg-{{ c.get(s,'secondary') }} text-capitalize">{{ s }}</span>{% endmacro %}
{% macro job_card(job) %}
 <div class="card h-100 shadow-sm card-hover"><div class="card-body">
  <h5 class="mb-1"><a class="stretched-link text-decoration-none" href="{{ url_for('job_detail', job_id=job.id) }}">{{ job.title }}</a></h5>
  <div class="text-muted small mb-2"><i class="bi bi-building"></i> {{ job.company }} &middot; <i class="bi bi-geo-alt"></i> {{ job.location }}</div>
  <span class="badge bg-light text-dark border">{{ job.job_type }}</span>
  {% if job.min_experience %}<span class="badge bg-light text-dark border">{{ job.min_experience }}+ yrs</span>{% endif %}
  {% if job.salary %}<span class="badge bg-light text-dark border">{{ job.salary }}</span>{% endif %}
  <div class="mt-2">{% for s in job.skill_list[:6] %}<span class="badge chip me-1">{{ s }}</span>{% endfor %}</div>
 </div></div>
{% endmacro %}""",

"index.html": """{% extends 'base.html' %}{% from 'macros.html' import job_card %}{% block content %}
<section class="hero p-5 mb-4"><div class="row align-items-center"><div class="col-lg-7">
 <h1 class="display-5 fw-bold">Find the right job. <span class="text-brand">Find the right talent.</span></h1>
 <p class="lead text-muted">An AI-powered job portal that screens resumes automatically, scores every candidate against the job requirements and recommends the best matches.</p>
 <form class="d-flex gap-2" action="{{ url_for('jobs') }}"><input class="form-control form-control-lg" name="q" placeholder="Job title, skill or company"><button class="btn btn-primary btn-lg"><i class="bi bi-search"></i></button></form>
</div><div class="col-lg-5 mt-4 mt-lg-0"><div class="row g-3">
 {% for label, value in [('Open jobs', stats.jobs), ('Companies', stats.companies), ('Candidates', stats.candidates), ('AI screenings', stats.applications)] %}
 <div class="col-6"><div class="card stat shadow-sm"><div class="card-body"><div class="text-muted small">{{ label }}</div><div class="fs-3 fw-bold">{{ value }}</div></div></div></div>{% endfor %}
</div></div></div></section>
<div class="row text-center mb-5">
 <div class="col-md-4"><i class="bi bi-file-earmark-text fs-1 text-brand"></i><h5>Resume Parsing</h5><p class="text-muted">PDF, DOCX or TXT: skills, experience and education are extracted automatically.</p></div>
 <div class="col-md-4"><i class="bi bi-graph-up-arrow fs-1 text-brand"></i><h5>AI Screening Score</h5><p class="text-muted">Skill matching + TF-IDF text similarity + experience fit = a 0-100 score.</p></div>
 <div class="col-md-4"><i class="bi bi-people fs-1 text-brand"></i><h5>Smart Matching</h5><p class="text-muted">Job recommendations for candidates, ranked applicants and talent search for employers.</p></div>
</div>
<h3>Latest jobs</h3><div class="row g-3">{% for job in jobs %}<div class="col-md-6 col-lg-4">{{ job_card(job) }}</div>{% else %}<p class="text-muted">No jobs yet.</p>{% endfor %}</div>
{% endblock %}""",

"login.html": """{% extends 'base.html' %}{% block content %}
<div class="row justify-content-center"><div class="col-md-5"><div class="card shadow-sm"><div class="card-body p-4">
 <h3>Login</h3><form method="post">""" + H + """
  <div class="mb-3"><label class="form-label">Email</label><input type="email" name="email" class="form-control" required autofocus></div>
  <div class="mb-3"><label class="form-label">Password</label><input type="password" name="password" class="form-control" required></div>
  <button class="btn btn-primary w-100">Login</button></form>
 <p class="mt-3 mb-0 small text-center">No account? <a href="{{ url_for('register') }}">Sign up</a></p>
</div></div>
<div class="alert alert-light border small mt-3"><b>Demo accounts</b> (password <code>password123</code>):<br>admin@jobmatch.ai &middot; hr@techsoft.com &middot; priya@example.com</div>
</div></div>{% endblock %}""",

"register.html": """{% extends 'base.html' %}{% block content %}
<div class="row justify-content-center"><div class="col-md-6"><div class="card shadow-sm"><div class="card-body p-4">
 <h3>Create an account</h3><form method="post">""" + H + """
  <div class="mb-3"><label class="form-label">I am a</label><select name="role" id="role" class="form-select" onchange="tc()">
   <option value="candidate">Job seeker (candidate)</option><option value="employer">Employer (recruiter)</option></select></div>
  <div class="mb-3"><label class="form-label">Full name</label><input name="name" class="form-control" required></div>
  <div class="mb-3"><label class="form-label">Email</label><input type="email" name="email" class="form-control" required></div>
  <div class="mb-3" id="co"><label class="form-label">Company name</label><input name="company" class="form-control"></div>
  <div class="mb-3"><label class="form-label">Password</label><input type="password" name="password" minlength="6" class="form-control" required></div>
  <button class="btn btn-primary w-100">Create account</button></form>
</div></div></div></div>
<script>function tc(){document.getElementById('co').style.display=document.getElementById('role').value=='employer'?'':'none'}tc()</script>
{% endblock %}""",

"jobs.html": """{% extends 'base.html' %}{% from 'macros.html' import job_card %}{% block content %}
<h2>Browse jobs</h2><form class="row g-2 mb-4">
 <div class="col-md-5"><input class="form-control" name="q" value="{{ q }}" placeholder="Title, skill, company"></div>
 <div class="col-md-3"><input class="form-control" name="location" value="{{ location }}" placeholder="Location"></div>
 <div class="col-md-2"><select class="form-select" name="job_type"><option value="">Any type</option>{% for t in job_types %}<option {% if t == job_type %}selected{% endif %}>{{ t }}</option>{% endfor %}</select></div>
 <div class="col-md-2"><button class="btn btn-primary w-100">Search</button></div></form>
<p class="text-muted">{{ jobs|length }} job(s) found</p>
<div class="row g-3">{% for job in jobs %}<div class="col-md-6 col-lg-4">{{ job_card(job) }}</div>{% else %}<p class="text-muted">No jobs match your search.</p>{% endfor %}</div>
{% endblock %}""",

"job_detail.html": """{% extends 'base.html' %}{% from 'macros.html' import badge, breakdown, chips, status %}{% block content %}
<div class="row g-4"><div class="col-lg-8"><div class="card shadow-sm"><div class="card-body p-4">
 {% if not job.is_open %}<span class="badge bg-secondary">Closed</span>{% endif %}
 <h2>{{ job.title }}</h2>
 <p class="text-muted">{{ job.company }} &middot; {{ job.location }} &middot; {{ job.job_type }}{% if job.salary %} &middot; {{ job.salary }}{% endif %} &middot; {{ job.min_experience }}+ years</p>
 <h5>Required skills</h5><div class="mb-3">{% for s in job.skill_list %}<span class="badge chip me-1">{{ s }}</span>{% endfor %}</div>
 <h5>Description</h5><p style="white-space:pre-wrap">{{ job.description }}</p>
</div></div></div>
<div class="col-lg-4">
 {% if not current_user.is_authenticated %}
  <div class="card shadow-sm"><div class="card-body"><p>Log in as a candidate to apply and see your AI match score.</p>
  <a class="btn btn-primary w-100" href="{{ url_for('login', next=request.path) }}">Login to apply</a></div></div>
 {% elif current_user.role == 'candidate' %}
  {% if existing %}
   <div class="card shadow-sm"><div class="card-body"><h5>You applied</h5><p>Status: {{ status(existing.status) }}</p>
   <p>AI score: {{ badge(existing.match_score) }} {{ existing.recommendation }}</p>
   <a class="btn btn-outline-primary w-100" href="{{ url_for('candidate_application', app_id=existing.id) }}">View screening report</a></div></div>
  {% elif job.is_open %}
   {% if preview %}<div class="card shadow-sm mb-3"><div class="card-body">
    <h6 class="text-muted">Your predicted match</h6><div class="big text-{{ preview.score|color }}">{{ '%.0f'|format(preview.score) }}%</div>
    <div class="mb-2">{{ preview.recommendation }}</div>{{ breakdown(preview.skill_score, preview.text_score, preview.experience_score) }}
    <div class="small"><b>Missing:</b> {{ chips(preview.missing, 'bg-danger') }}</div></div></div>{% endif %}
   <div class="card shadow-sm"><div class="card-body"><h5>Apply now</h5>
    <form method="post" action="{{ url_for('apply', job_id=job.id) }}" enctype="multipart/form-data">""" + H + """
     <div class="mb-3"><label class="form-label">Resume (PDF, DOCX, TXT)</label>
      <input type="file" name="resume" class="form-control" accept=".pdf,.docx,.txt" {% if not current_user.resume_text %}required{% endif %}>
      {% if current_user.resume_text %}<div class="form-text">Leave empty to use your profile resume.</div>{% endif %}</div>
     <div class="mb-3"><label class="form-label">Cover letter (optional)</label><textarea name="cover_letter" rows="3" class="form-control"></textarea></div>
     <button class="btn btn-success w-100">Submit application</button></form></div></div>
  {% else %}<div class="alert alert-secondary">This job is closed.</div>{% endif %}
 {% elif current_user.id == job.employer_id or current_user.role == 'admin' %}
  <div class="card shadow-sm"><div class="card-body d-grid gap-2">
   <a class="btn btn-primary" href="{{ url_for('applicants', job_id=job.id) }}">Ranked applicants ({{ job.applications|length }})</a>
   <a class="btn btn-outline-primary" href="{{ url_for('talent_search', job_id=job.id) }}">AI talent search</a>
   <a class="btn btn-outline-secondary" href="{{ url_for('edit_job', job_id=job.id) }}">Edit job</a></div></div>
 {% endif %}
</div></div>{% endblock %}""",

"candidate_dashboard.html": """{% extends 'base.html' %}{% from 'macros.html' import badge, status %}{% block content %}
<h2>Hi, {{ current_user.name }}</h2>
{% if not current_user.resume_text %}<div class="alert alert-info"><a href="{{ url_for('candidate_profile') }}">Upload your resume</a> to get AI job recommendations.</div>{% endif %}
<div class="row g-4"><div class="col-lg-7"><h4>My applications</h4><div class="card shadow-sm"><div class="table-responsive"><table class="table align-middle mb-0">
 <thead class="table-light"><tr><th>Job</th><th>AI score</th><th>Status</th><th></th></tr></thead><tbody>
 {% for a in applications %}<tr><td><a href="{{ url_for('job_detail', job_id=a.job_id) }}">{{ a.job.title }}</a><div class="small text-muted">{{ a.job.company }}</div></td>
  <td>{{ badge(a.match_score) }}</td><td>{{ status(a.status) }}</td>
  <td><a class="btn btn-sm btn-outline-primary" href="{{ url_for('candidate_application', app_id=a.id) }}">Report</a></td></tr>
 {% else %}<tr><td colspan="4" class="text-center text-muted py-4">No applications yet. <a href="{{ url_for('jobs') }}">Browse jobs</a></td></tr>{% endfor %}
</tbody></table></div></div></div>
<div class="col-lg-5"><h4><i class="bi bi-stars text-brand"></i> Recommended for you</h4>
 {% for job, r in recommendations %}<div class="card shadow-sm mb-2 card-hover"><div class="card-body py-2"><div class="d-flex justify-content-between align-items-center">
  <div><a class="fw-semibold stretched-link text-decoration-none" href="{{ url_for('job_detail', job_id=job.id) }}">{{ job.title }}</a><div class="small text-muted">{{ job.company }} &middot; {{ job.location }}</div></div>{{ badge(r.score) }}</div>
  {% if r.missing %}<div class="small text-muted">Skills to learn: {{ r.missing[:4]|join(', ') }}</div>{% endif %}</div></div>
 {% else %}<p class="text-muted">Add skills or a resume to your profile to get recommendations.</p>{% endfor %}
</div></div>{% endblock %}""",

"candidate_profile.html": """{% extends 'base.html' %}{% from 'macros.html' import chips %}{% block content %}
<div class="row g-4"><div class="col-lg-6"><div class="card shadow-sm"><div class="card-body p-4"><h3>My profile</h3>
 <form method="post" enctype="multipart/form-data">""" + H + """
  <div class="mb-3"><label class="form-label">Name</label><input name="name" class="form-control" value="{{ current_user.name }}"></div>
  <div class="mb-3"><label class="form-label">Headline</label><input name="headline" class="form-control" value="{{ current_user.headline or '' }}" placeholder="Final-year B.Tech CSE student"></div>
  <div class="mb-3"><label class="form-label">Location</label><input name="location" class="form-control" value="{{ current_user.location or '' }}"></div>
  <div class="mb-3"><label class="form-label">Skills (comma-separated)</label><input name="skills" class="form-control" value="{{ current_user.skills or '' }}"></div>
  <div class="mb-3"><label class="form-label">Resume (PDF, DOCX, TXT - max 5 MB)</label><input type="file" name="resume" class="form-control" accept=".pdf,.docx,.txt"></div>
  <button class="btn btn-primary">Save profile</button></form></div></div></div>
<div class="col-lg-6"><div class="card shadow-sm"><div class="card-body p-4"><h4><i class="bi bi-cpu text-brand"></i> AI resume analysis</h4>
 {% if analysis %}
  <p class="mb-1"><b>Experience:</b> {{ ('%g years'|format(analysis.years)) if analysis.years is not none else 'Not mentioned' }}</p>
  <p class="mb-1"><b>Education:</b> {{ analysis.education or 'Not detected' }}</p>
  <p><b>Word count:</b> {{ analysis.words }}</p>
  <h6>Skills found in your resume ({{ analysis.skills|length }})</h6>{{ chips(analysis.skills, 'chip') }}
 {% else %}<p class="text-muted">Upload a resume to see the skills, experience and education the AI extracts.</p>{% endif %}
</div></div></div></div>{% endblock %}""",

"report.html": """{% extends 'base.html' %}{% from 'macros.html' import breakdown, chips, status %}{% block content %}
<h2>{{ a.candidate.name if employer_view else a.job.title }} <small class="text-muted fs-6">{{ 'for ' ~ a.job.title if employer_view else 'at ' ~ a.job.company }}</small></h2>
<p>Status: {{ status(a.status) }} &middot; Applied {{ a.created_at.strftime('%d %b %Y') }}</p>
<div class="row g-4"><div class="col-lg-4">
 <div class="card shadow-sm mb-3"><div class="card-body text-center"><div class="text-muted small">AI match score</div>
  <div class="big text-{{ a.match_score|color }}">{{ '%.0f'|format(a.match_score) }}%</div><div class="fw-semibold">{{ a.recommendation }}</div><hr>
  {{ breakdown(a.skill_score, a.text_score, a.experience_score) }}</div></div>
 <div class="card shadow-sm"><div class="card-body">
  <p class="mb-1"><b>Experience:</b> {{ ('%g years'|format(a.years_experience)) if a.years_experience is not none else 'Not mentioned' }}</p>
  <p class="mb-2"><b>Education:</b> {{ a.education or 'Not detected' }}</p>
  {% if employer_view %}<p class="mb-2"><b>Email:</b> {{ a.candidate.email }}</p>
   <form method="post" action="{{ url_for('update_status', app_id=a.id) }}" class="d-flex gap-2">""" + H + """
   <select name="status" class="form-select">{% for s in statuses %}<option value="{{ s }}" {% if s == a.status %}selected{% endif %}>{{ s|capitalize }}</option>{% endfor %}</select>
   <button class="btn btn-primary">Update</button></form>
  {% elif a.status == 'applied' %}
   <form method="post" action="{{ url_for('withdraw', app_id=a.id) }}" onsubmit="return confirm('Withdraw?')">""" + H + """<button class="btn btn-outline-danger w-100">Withdraw application</button></form>
  {% endif %}
  {% if a.resume_filename %}<a class="btn btn-outline-secondary w-100 mt-2" href="{{ url_for('download_resume', app_id=a.id) }}">Download resume</a>{% endif %}
 </div></div></div>
<div class="col-lg-8"><div class="card shadow-sm mb-3"><div class="card-body">
 <h5>Matched skills</h5><div class="mb-3">{{ chips(a.matched_list) }}</div>
 <h5>{{ 'Missing skills' if employer_view else 'Skills to improve' }}</h5>{{ chips(a.missing_list, 'bg-danger') }}</div></div>
 {% if employer_view %}<div class="card shadow-sm"><div class="card-body"><h5>Extracted resume text</h5><div class="resume p-3 rounded">{{ a.resume_text }}</div></div></div>{% endif %}
</div></div>{% endblock %}""",

"employer_dashboard.html": """{% extends 'base.html' %}{% block content %}
<div class="d-flex justify-content-between align-items-center mb-3"><h2>{{ current_user.company }} - Dashboard</h2><a class="btn btn-primary" href="{{ url_for('new_job') }}">+ Post a job</a></div>
<div class="row g-3 mb-4">{% for label, v in [('Job postings', jobs|length), ('Applications received', total), ('Shortlisted', shortlisted)] %}
 <div class="col-md-4"><div class="card stat shadow-sm"><div class="card-body"><div class="text-muted small">{{ label }}</div><div class="fs-3 fw-bold">{{ v }}</div></div></div></div>{% endfor %}</div>
<div class="card shadow-sm"><div class="table-responsive"><table class="table table-hover align-middle mb-0">
 <thead class="table-light"><tr><th>Job</th><th>Status</th><th>Applicants</th><th class="text-end">Actions</th></tr></thead><tbody>
 {% for job in jobs %}<tr><td><a href="{{ url_for('job_detail', job_id=job.id) }}">{{ job.title }}</a><div class="small text-muted">{{ job.location }}</div></td>
  <td>{% if job.is_open %}<span class="badge bg-success">Open</span>{% else %}<span class="badge bg-secondary">Closed</span>{% endif %}</td>
  <td>{{ job.applications|length }}</td><td class="text-end text-nowrap">
  <a class="btn btn-sm btn-primary" href="{{ url_for('applicants', job_id=job.id) }}">Applicants</a>
  <a class="btn btn-sm btn-outline-primary" href="{{ url_for('talent_search', job_id=job.id) }}">Talent search</a>
  <a class="btn btn-sm btn-outline-secondary" href="{{ url_for('edit_job', job_id=job.id) }}">Edit</a>
  <form class="d-inline" method="post" action="{{ url_for('toggle_job', job_id=job.id) }}">""" + H + """<button class="btn btn-sm btn-outline-warning">{{ 'Close' if job.is_open else 'Reopen' }}</button></form>
  <form class="d-inline" method="post" action="{{ url_for('delete_job', job_id=job.id) }}" onsubmit="return confirm('Delete this job?')">""" + H + """<button class="btn btn-sm btn-outline-danger">Delete</button></form></td></tr>
 {% else %}<tr><td colspan="4" class="text-center text-muted py-4">No jobs yet.</td></tr>{% endfor %}
</tbody></table></div></div>{% endblock %}""",

"job_form.html": """{% extends 'base.html' %}{% block content %}
<div class="row justify-content-center"><div class="col-lg-8"><div class="card shadow-sm"><div class="card-body p-4">
 <h3>{{ 'Edit job' if editing else 'Post a new job' }}</h3><form method="post">""" + H + """<div class="row g-3">
  <div class="col-md-8"><label class="form-label">Job title *</label><input name="title" class="form-control" value="{{ job.title or '' }}" required></div>
  <div class="col-md-4"><label class="form-label">Company *</label><input name="company" class="form-control" value="{{ job.company or '' }}" required></div>
  <div class="col-md-4"><label class="form-label">Location *</label><input name="location" class="form-control" value="{{ job.location or '' }}" required></div>
  <div class="col-md-3"><label class="form-label">Type</label><select name="job_type" class="form-select">{% for t in job_types %}<option {% if t == job.job_type %}selected{% endif %}>{{ t }}</option>{% endfor %}</select></div>
  <div class="col-md-2"><label class="form-label">Min exp (yrs)</label><input type="number" min="0" name="min_experience" class="form-control" value="{{ job.min_experience or 0 }}"></div>
  <div class="col-md-3"><label class="form-label">Salary</label><input name="salary" class="form-control" value="{{ job.salary or '' }}"></div>
  <div class="col-12"><label class="form-label">Required skills * (comma-separated, used by the AI)</label><input name="required_skills" class="form-control" value="{{ job.required_skills or '' }}" placeholder="python, flask, sql" required></div>
  <div class="col-12"><label class="form-label">Description *</label><textarea name="description" rows="7" class="form-control" required>{{ job.description or '' }}</textarea></div>
 </div><button class="btn btn-primary mt-3">{{ 'Save & re-screen applicants' if editing else 'Post job' }}</button></form>
</div></div></div></div>{% endblock %}""",

"applicants.html": """{% extends 'base.html' %}{% from 'macros.html' import badge, chips %}{% block content %}
<div class="d-flex justify-content-between flex-wrap gap-2 mb-3"><div><h2 class="mb-0">Applicants ranked by AI</h2><div class="text-muted">{{ job.title }} &middot; {{ job.company }}</div></div>
 <div class="d-flex gap-2"><a class="btn btn-outline-primary" href="{{ url_for('talent_search', job_id=job.id) }}">AI talent search</a>
 <form method="post" action="{{ url_for('auto_shortlist', job_id=job.id) }}" class="d-flex gap-1">""" + H + """
  <input type="number" name="threshold" value="70" class="form-control" style="width:5rem"><button class="btn btn-success text-nowrap">Auto-shortlist</button></form></div></div>
<form class="row g-2 mb-3"><div class="col-auto"><select name="status" class="form-select"><option value="">All statuses</option>{% for s in statuses %}<option value="{{ s }}" {% if s == status %}selected{% endif %}>{{ s|capitalize }}</option>{% endfor %}</select></div>
 <div class="col-auto"><input type="number" name="min_score" class="form-control" placeholder="Min score" value="{{ min_score or '' }}"></div><div class="col-auto"><button class="btn btn-secondary">Filter</button></div></form>
<div class="card shadow-sm"><div class="table-responsive"><table class="table align-middle mb-0">
 <thead class="table-light"><tr><th>#</th><th>Candidate</th><th>AI score</th><th>Matched / missing skills</th><th>Exp.</th><th>Education</th><th>Status</th><th></th></tr></thead><tbody>
 {% for a in applications %}<tr><td class="fw-bold">{{ loop.index }}</td><td>{{ a.candidate.name }}<div class="small text-muted">{{ a.candidate.email }}</div></td>
  <td>{{ badge(a.match_score) }}<div class="small text-muted">{{ a.recommendation }}</div></td>
  <td style="max-width:320px">{{ chips(a.matched_list) }}{% for s in a.missing_list %}<span class="badge bg-danger-subtle text-danger-emphasis me-1">{{ s }}</span>{% endfor %}</td>
  <td>{{ ('%g yrs'|format(a.years_experience)) if a.years_experience is not none else '-' }}</td><td>{{ a.education or '-' }}</td>
  <td><form method="post" action="{{ url_for('update_status', app_id=a.id) }}">""" + H + """<select name="status" class="form-select form-select-sm" onchange="this.form.submit()">
   {% for s in statuses %}<option value="{{ s }}" {% if s == a.status %}selected{% endif %}>{{ s|capitalize }}</option>{% endfor %}</select></form></td>
  <td><a class="btn btn-sm btn-outline-primary" href="{{ url_for('employer_application', app_id=a.id) }}">Report</a></td></tr>
 {% else %}<tr><td colspan="8" class="text-center text-muted py-4">No applicants yet.</td></tr>{% endfor %}
</tbody></table></div></div>{% endblock %}""",

"talent.html": """{% extends 'base.html' %}{% from 'macros.html' import badge, chips %}{% block content %}
<h2><i class="bi bi-stars text-brand"></i> AI talent search</h2>
<p class="text-muted">Every candidate on the platform ranked against <b>{{ job.title }}</b>, including people who haven't applied yet.</p>
<div class="card shadow-sm"><div class="table-responsive"><table class="table align-middle mb-0">
 <thead class="table-light"><tr><th>#</th><th>Candidate</th><th>AI score</th><th>Matched</th><th>Missing</th><th></th></tr></thead><tbody>
 {% for c, r in ranked %}<tr><td class="fw-bold">{{ loop.index }}</td><td>{{ c.name }}<div class="small text-muted">{{ c.headline or '' }}</div></td>
  <td>{{ badge(r.score) }}<div class="small text-muted">{{ r.recommendation }}</div></td><td>{{ chips(r.matched) }}</td><td>{{ chips(r.missing, 'bg-danger') }}</td>
  <td>{% if c.id in applied %}<span class="badge bg-info">Applied</span>{% else %}<a class="btn btn-sm btn-outline-primary" href="mailto:{{ c.email }}?subject=Opportunity: {{ job.title }}">Invite</a>{% endif %}</td></tr>
 {% else %}<tr><td colspan="6" class="text-center text-muted py-4">No candidate profiles yet.</td></tr>{% endfor %}
</tbody></table></div></div>{% endblock %}""",

"admin_dashboard.html": """{% extends 'base.html' %}{% from 'macros.html' import badge, status %}{% block content %}
<h2>Admin dashboard</h2><div class="row g-3 mb-4">
 {% for label, v in [('Candidates', stats.candidates), ('Employers', stats.employers), ('Jobs', stats.jobs), ('Applications', stats.applications), ('Avg. AI score', '%.0f%%'|format(stats.avg))] %}
 <div class="col"><div class="card stat shadow-sm"><div class="card-body"><div class="text-muted small">{{ label }}</div><div class="fs-4 fw-bold">{{ v }}</div></div></div></div>{% endfor %}</div>
<div class="row g-4 mb-4"><div class="col-md-6"><div class="card shadow-sm h-100"><div class="card-body"><h5>AI score distribution</h5>
 {% for b, n in buckets.items() %}<div class="d-flex align-items-center mb-2"><span style="width:4.5rem">{{ b }}</span><div class="progress flex-grow-1" style="height:18px">
  <div class="progress-bar bg-{{ b.split('-')[0]|int|color }}" style="width:{{ n / (stats.applications or 1) * 100 }}%">{{ n }}</div></div></div>{% endfor %}</div></div></div>
 <div class="col-md-6"><div class="card shadow-sm h-100"><div class="card-body"><h5>Applications by status</h5>
 {% for s, n in by_status.items() %}<div class="d-flex justify-content-between border-bottom py-1">{{ status(s) }}<b>{{ n }}</b></div>{% endfor %}</div></div></div></div>
<h5>Recent applications</h5><div class="card shadow-sm"><table class="table mb-0"><thead class="table-light"><tr><th>Candidate</th><th>Job</th><th>Score</th><th>Status</th></tr></thead><tbody>
 {% for a in recent %}<tr><td>{{ a.candidate.name }}</td><td>{{ a.job.title }}</td><td>{{ badge(a.match_score) }}</td><td>{{ status(a.status) }}</td></tr>{% endfor %}</tbody></table></div>
{% endblock %}""",

"admin_users.html": """{% extends 'base.html' %}{% block content %}
<h2>Manage users</h2><form class="row g-2 mb-3"><div class="col-md-4"><input name="q" class="form-control" value="{{ q }}" placeholder="Search name or email"></div>
 <div class="col-md-3"><select name="role" class="form-select"><option value="">All roles</option>{% for r in ['candidate','employer','admin'] %}<option value="{{ r }}" {% if r == role %}selected{% endif %}>{{ r|capitalize }}</option>{% endfor %}</select></div>
 <div class="col-auto"><button class="btn btn-secondary">Filter</button></div></form>
<div class="card shadow-sm"><div class="table-responsive"><table class="table align-middle mb-0"><thead class="table-light"><tr><th>Name</th><th>Email</th><th>Role</th><th>Status</th><th></th></tr></thead><tbody>
 {% for u in users %}<tr><td>{{ u.name }}</td><td>{{ u.email }}</td><td class="text-capitalize">{{ u.role }}</td>
  <td>{% if u.active_account %}<span class="badge bg-success">Active</span>{% else %}<span class="badge bg-secondary">Deactivated</span>{% endif %}</td>
  <td class="text-end text-nowrap">{% if u.id != current_user.id %}
   <form class="d-inline" method="post" action="{{ url_for('toggle_user', user_id=u.id) }}">""" + H + """<button class="btn btn-sm btn-outline-warning">{{ 'Deactivate' if u.active_account else 'Activate' }}</button></form>
   <form class="d-inline" method="post" action="{{ url_for('delete_user', user_id=u.id) }}" onsubmit="return confirm('Delete user and all their data?')">""" + H + """<button class="btn btn-sm btn-outline-danger">Delete</button></form>{% endif %}</td></tr>{% endfor %}
</tbody></table></div></div>{% endblock %}""",

"admin_jobs.html": """{% extends 'base.html' %}{% block content %}
<h2>All jobs</h2><div class="card shadow-sm"><div class="table-responsive"><table class="table align-middle mb-0"><thead class="table-light"><tr><th>Job</th><th>Employer</th><th>Applicants</th><th>Status</th><th></th></tr></thead><tbody>
 {% for job in jobs %}<tr><td><a href="{{ url_for('job_detail', job_id=job.id) }}">{{ job.title }}</a><div class="small text-muted">{{ job.company }}</div></td><td>{{ job.employer.name }}</td>
  <td><a href="{{ url_for('applicants', job_id=job.id) }}">{{ job.applications|length }}</a></td>
  <td>{% if job.is_open %}<span class="badge bg-success">Open</span>{% else %}<span class="badge bg-secondary">Closed</span>{% endif %}</td>
  <td class="text-end text-nowrap"><form class="d-inline" method="post" action="{{ url_for('toggle_job', job_id=job.id) }}">""" + H + """<button class="btn btn-sm btn-outline-warning">{{ 'Close' if job.is_open else 'Reopen' }}</button></form>
  <form class="d-inline" method="post" action="{{ url_for('delete_job', job_id=job.id) }}" onsubmit="return confirm('Delete?')">""" + H + """<button class="btn btn-sm btn-outline-danger">Delete</button></form></td></tr>{% endfor %}
</tbody></table></div></div>{% endblock %}""",

"error.html": """{% extends 'base.html' %}{% block content %}<div class="text-center py-5"><div class="display-1 fw-bold text-brand">{{ code }}</div><p class="lead">{{ message }}</p><a class="btn btn-primary" href="{{ url_for('index') }}">Go home</a></div>{% endblock %}""",
}

app.jinja_loader = DictLoader(TEMPLATES)


# =============================================================================
# 12. DEMO DATA (created automatically on first run)
# =============================================================================
DEMO_RESUMES = {
    "priya": """PRIYA SHARMA - Data Scientist, Chennai
Data scientist with 3 years of experience building machine learning models and NLP pipelines.
EDUCATION: M.Tech Computer Science (AI), SRM Institute of Science and Technology
SKILLS: Python, SQL, Machine Learning, Deep Learning, NLP, scikit-learn, TensorFlow, PyTorch,
Pandas, NumPy, Matplotlib, Tableau, Statistics, Git, Docker, AWS, Flask
EXPERIENCE: Data Scientist, ShopKart Analytics - built a recommendation engine and a sentiment
analysis NLP model; deployed models as REST APIs with Flask and Docker on AWS.""",
    "rahul": """RAHUL VERMA - Full Stack Developer, Bengaluru
Full stack developer with 4 years of experience building scalable web applications.
EDUCATION: B.E. Information Technology
SKILLS: JavaScript, TypeScript, ReactJS, Node.js, Express, HTML5, CSS3, Bootstrap, Python, Django,
REST APIs, GraphQL, MongoDB, PostgreSQL, MySQL, Docker, Kubernetes, AWS, CI/CD, Git, Agile
EXPERIENCE: Senior Software Engineer, CloudNova - led a React + Node.js SaaS dashboard.""",
    "ananya": """ANANYA IYER - Final Year B.Tech CSE Student, SRM University
Motivated fresher seeking an entry-level software or data analyst role.
EDUCATION: B.Tech Computer Science and Engineering, SRM IST (2022-2026), CGPA 8.9
SKILLS: Python, Java, C, C++, HTML, CSS, JavaScript, SQL, MySQL, Excel, Power BI, Git
Soft skills: Communication, Teamwork, Problem Solving
PROJECTS: Attendance system (Python, Flask, SQLite); crop-yield data analysis (Pandas, Matplotlib).
INTERNSHIP: Data Analyst Intern, InfoMetrics - Excel and Power BI dashboards.""",
    "arjun": """ARJUN NAIR - DevOps / Cloud Engineer, Hyderabad
5+ years of experience in cloud infrastructure, automation and site reliability.
SKILLS: AWS, Azure, GCP, Terraform, Docker, Kubernetes, Jenkins, CI/CD, Linux, Bash, Python, Git
EXPERIENCE: Cloud Engineer, InfraScale - multi-region AWS with Terraform; CI/CD for 40 microservices.
EDUCATION: B.Tech Electronics and Communication""",
}

DEMO_JOBS = [
    ("Machine Learning Engineer", "Chennai", "Full-time", "Rs 12-18 LPA", 2,
     "python, machine learning, deep learning, nlp, tensorflow, scikit-learn, sql, docker",
     "Design, train and deploy machine learning and NLP models for product search and recommendations. "
     "Build data pipelines and serve models via REST APIs."),
    ("Full Stack Developer (React + Node)", "Bengaluru", "Full-time", "Rs 10-16 LPA", 3,
     "javascript, react, node.js, express, mongodb, rest api, git, docker",
     "Build web applications end-to-end with React frontends, Node.js/Express backends and MongoDB."),
    ("Data Analyst - Graduate Trainee", "Chennai", "Full-time", "Rs 4-6 LPA", 0,
     "sql, excel, power bi, python, data analysis, communication",
     "Entry-level role for fresh graduates. Analyse business data, build Power BI and Excel dashboards, "
     "write SQL queries and present insights."),
    ("DevOps Engineer", "Hyderabad", "Full-time", "Rs 14-22 LPA", 4,
     "aws, docker, kubernetes, terraform, jenkins, ci/cd, linux, python",
     "Own our AWS cloud infrastructure, automate with Terraform, run Kubernetes and build CI/CD pipelines."),
    ("Python Backend Developer Intern", "Remote", "Internship", "Rs 25,000 / month", 0,
     "python, flask, sql, rest api, git",
     "6-month internship building REST APIs with Python and Flask, writing SQL and unit tests."),
]


def seed_demo_data():
    def user(name, email, role, **kw):
        u = User(name=name, email=email, role=role, **kw)
        u.set_password("password123")
        db.session.add(u)
        return u

    user("Platform Admin", "admin@jobmatch.ai", "admin")
    emp1 = user("Kavya Reddy", "hr@techsoft.com", "employer", company="TechSoft Solutions")
    emp2 = user("Vikram Singh", "talent@cloudnine.io", "employer", company="CloudNine Systems")
    cands = [
        user("Priya Sharma", "priya@example.com", "candidate", headline="Data Scientist", location="Chennai",
             skills="python, machine learning, nlp", resume_text=DEMO_RESUMES["priya"]),
        user("Rahul Verma", "rahul@example.com", "candidate", headline="Full Stack Developer", location="Bengaluru",
             skills="react, node.js, python", resume_text=DEMO_RESUMES["rahul"]),
        user("Ananya Iyer", "ananya@example.com", "candidate", headline="B.Tech CSE final-year student",
             location="Chennai", skills="python, sql, excel", resume_text=DEMO_RESUMES["ananya"]),
        user("Arjun Nair", "arjun@example.com", "candidate", headline="DevOps Engineer", location="Hyderabad",
             skills="aws, kubernetes, terraform", resume_text=DEMO_RESUMES["arjun"]),
    ]
    jobs_ = []
    for i, (title, loc, jt, sal, exp, skills, desc) in enumerate(DEMO_JOBS):
        emp = emp1 if i % 2 == 0 else emp2
        job = Job(title=title, company=emp.company, location=loc, job_type=jt, salary=sal,
                  min_experience=exp, required_skills=skills, description=desc, employer=emp)
        db.session.add(job)
        jobs_.append(job)
    db.session.flush()
    for ci, job_ids in {0: [0, 2, 4], 1: [1, 4, 0], 2: [2, 4, 0, 1], 3: [3, 1]}.items():
        c = cands[ci]
        for j in job_ids:
            a = Application(job=jobs_[j], candidate=c, resume_text=c.resume_text,
                            cover_letter="I am excited to apply for this role.")
            a.apply_screening(score_resume(c.resume_text, jobs_[j], c.skill_list))
            db.session.add(a)
    db.session.commit()
    print("Demo data created. Login password for all demo accounts: password123")


with app.app_context():
    db.create_all()
    if not app.config.get("TESTING") and User.query.count() == 0:
        seed_demo_data()


# =============================================================================
# 13. RUN THE APP
# =============================================================================
if __name__ == "__main__":
    print("JobMatch AI running at  http://127.0.0.1:5000")
    app.run(debug=True)
