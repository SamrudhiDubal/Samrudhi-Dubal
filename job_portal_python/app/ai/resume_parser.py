"""Extract plain text from uploaded resumes (PDF, DOCX, TXT)."""
import io
import os


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return _from_pdf(file_bytes)
    if ext == ".docx":
        return _from_docx(file_bytes)
    if ext == ".txt":
        return file_bytes.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported resume format: {ext}")


def _from_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _from_docx(file_bytes: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
